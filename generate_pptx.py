import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path="ActionPilot_Pitch_Deck.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]  # completely blank

    # Color Palette (Dark Theme matching ActionPilot)
    BG_COLOR = RGBColor(11, 15, 25)          # #0B0F19
    CARD_BG = RGBColor(18, 24, 38)          # #121826
    CARD_BORDER = RGBColor(40, 50, 75)      # Slate border
    PRIMARY = RGBColor(59, 130, 246)        # Electric Blue #3B82F6
    ACCENT = RGBColor(16, 185, 129)         # Emerald Green #10B981
    AMBER = RGBColor(245, 158, 11)          # Amber #F59E0B
    DANGER = RGBColor(239, 68, 68)          # Crimson #EF4444
    TEXT_MAIN = RGBColor(248, 250, 252)     # White #F8FAFC
    TEXT_MUTED = RGBColor(148, 163, 184)    # Gray #94A3B8

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, tag_text, title_text, subtitle_text=""):
        # Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(0.4))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = tag_text.upper()
        p_tag.font.bold = True
        p_tag.font.size = Pt(11)
        p_tag.font.color.rgb = PRIMARY

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.9))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.bold = True
        p_title.font.size = Pt(28)
        p_title.font.color.rgb = TEXT_MAIN

        # Subtitle
        if subtitle_text:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.65), Inches(11.5), Inches(0.6))
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitle_text
            p_sub.font.size = Pt(14)
            p_sub.font.color.rgb = TEXT_MUTED

    def add_card(slide, x, y, w, h, title, body_lines, border_color=CARD_BORDER, icon=""):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)

        tb = slide.shapes.add_textbox(x + Inches(0.2), y + Inches(0.2), w - Inches(0.4), h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = f"{icon}  {title}" if icon else title
        p_t.font.bold = True
        p_t.font.size = Pt(17)
        p_t.font.color.rgb = TEXT_MAIN

        for line in body_lines:
            p_b = tf.add_paragraph()
            p_b.text = line
            p_b.font.size = Pt(13)
            p_b.font.color.rgb = TEXT_MUTED
            p_b.space_before = Pt(8)

        return card

    # ==================== SLIDE 1: COVER ====================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1)

    # Logo / Badge
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(2.8), Inches(0.4))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(25, 45, 80)
    badge.line.color.rgb = PRIMARY
    b_tf = badge.text_frame
    b_tf.paragraphs[0].text = "⚡ HACKATHON 2026 EDITION"
    b_tf.paragraphs[0].font.bold = True
    b_tf.paragraphs[0].font.size = Pt(10)
    b_tf.paragraphs[0].font.color.rgb = PRIMARY
    b_tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Main Title
    tb_title = s1.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.5), Inches(2.0))
    tf1 = tb_title.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "ActionPilot"
    p1.font.bold = True
    p1.font.size = Pt(54)
    p1.font.color.rgb = TEXT_MAIN

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Evidence-Driven AI Workflow Automation & Governance"
    p1_sub.font.bold = True
    p1_sub.font.size = Pt(24)
    p1_sub.font.color.rgb = PRIMARY
    p1_sub.space_before = Pt(10)

    p1_desc = tf1.add_paragraph()
    p1_desc.text = "Turning messy human meetings, voice memos, and incident chats into deterministic, verifiable, and auditable task execution with zero hallucinations."
    p1_desc.font.size = Pt(15)
    p1_desc.font.color.rgb = TEXT_MUTED
    p1_desc.space_before = Pt(14)

    # 3 Pillars
    add_card(s1, Inches(0.8), Inches(4.3), Inches(3.6), Inches(2.4),
             "100% Provenance",
             ["Every extracted task is strictly grounded to character-level source offsets.", "Zero hallucinations guaranteed."],
             PRIMARY, "🔍")
    add_card(s1, Inches(4.8), Inches(4.3), Inches(3.6), Inches(2.4),
             "AI Conflict Resolver",
             ["Detects contradictory deadlines between teams and prompts human resolution.", "No silent failures."],
             AMBER, "⚔️")
    add_card(s1, Inches(8.8), Inches(4.3), Inches(3.6), Inches(2.4),
             "1-Click Dispatch Hub",
             ["Instantly triggers Google Calendar, Slack webhooks, GitHub issues, and Email.", "Includes dry-run pitch mode."],
             ACCENT, "🚀")

    notes1 = s1.notes_slide.notes_text_frame
    notes1.text = (
        "HOOK (15s):\n"
        "Judges, every company loses hours turning messy meetings, voice memos, and emails into Jira, Slack, "
        "and Calendar tasks. But trusting standard AI creates hallucinations and costly mistakes. "
        "We built ActionPilot — an evidence-grounded AI orchestration engine with zero hallucination, "
        "conflict resolution, and human governance."
    )

    # ==================== SLIDE 2: THE PROBLEM ====================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "Market Pain Point", "The 3 Deadly Flaws of AI Workflow Tools",
               "Why autonomous agents and LLM tools fail in high-stakes enterprise environments:")

    add_card(s2, Inches(0.8), Inches(2.5), Inches(3.6), Inches(4.2),
             "1. Hallucination Risk",
             ["• Unconstrained LLMs invent dates, assignees, and tasks when context is ambiguous.",
              "• No reference back to the original sentence or speaker.",
              "• Creates costly errors and destroys trust in enterprise operations."],
             DANGER, "🚨")

    add_card(s2, Inches(4.8), Inches(2.5), Inches(3.6), Inches(4.2),
             "2. Contradiction Blindness",
             ["• Cross-functional meetings frequently contain conflicting deadlines.",
              "• E.g., Product Lead says 'Launch Nov 2', DevOps mandates 'Freeze Oct 20'.",
              "• Current tools silently pick the wrong date or crash."],
             AMBER, "⚠️")

    add_card(s2, Inches(8.8), Inches(2.5), Inches(3.6), Inches(4.2),
             "3. Missing Governance",
             ["• Black-box autonomous agents trigger real-world side effects without sign-off.",
              "• No cryptographic audit trail of who approved what.",
              "• Violates enterprise SOC-2, HIPAA, and security compliance policies."],
             PRIMARY, "🔒")

    notes2 = s2.notes_slide.notes_text_frame
    notes2.text = (
        "THE PROBLEM (30s):\n"
        "Why can't enterprises trust current AI agents? Three deadly flaws:\n"
        "1. Hallucination Risk: Blind LLMs invent deadlines and assignees.\n"
        "2. Contradiction Blindness: When two leaders give conflicting dates in separate meetings, LLMs guess or break.\n"
        "3. Lack of Governance: Autonomous tools execute actions without immutable audit trails."
    )

    # ==================== SLIDE 3: ARCHITECTURE ====================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "System Architecture", "How ActionPilot Solves the Execution Chain",
               "A 5-step deterministic, audit-proof pipeline from raw voice/text to external execution:")

    # 5 Pipeline cards
    steps = [
        ("01. INGEST", "Web Speech API / Voice memos / PDFs / Meeting transcripts", PRIMARY, "🎙️"),
        ("02. GROUND", "Extract actions linked to exact character offsets & quotes", PRIMARY, "🎯"),
        ("03. CONFLICT", "Detect contradictory dates & trigger human override", AMBER, "⚔️"),
        ("04. SIGN-OFF", "Human approval gate with immutable audit logging", ACCENT, "🛡️"),
        ("05. DISPATCH", "Direct sync to Google Calendar, Slack, GitHub, Email", ACCENT, "🚀"),
    ]

    for idx, (title, desc, col, icon) in enumerate(steps):
        x = Inches(0.8 + idx * 2.38)
        c = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.6), Inches(2.2), Inches(2.5))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = col
        c.line.width = Pt(1.5)
        ctf = c.text_frame
        ctf.word_wrap = True
        p0 = ctf.paragraphs[0]
        p0.text = f"{icon} {title}"
        p0.font.bold = True
        p0.font.size = Pt(13)
        p0.font.color.rgb = col
        p1 = ctf.add_paragraph()
        p1.text = desc
        p1.font.size = Pt(11)
        p1.font.color.rgb = TEXT_MUTED
        p1.space_before = Pt(8)

    # Code / Provenance box at bottom
    code_box = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(5.4), Inches(11.7), Inches(1.4))
    code_box.fill.solid()
    code_box.fill.fore_color.rgb = RGBColor(15, 20, 30)
    code_box.line.color.rgb = RGBColor(40, 50, 75)
    cbtf = code_box.text_frame
    cbtf.word_wrap = True
    cp = cbtf.paragraphs[0]
    cp.text = "// Grounded Provenance Engine Output:"
    cp.font.size = Pt(11)
    cp.font.color.rgb = PRIMARY
    cp2 = cbtf.add_paragraph()
    cp2.text = 'Task: "Prepare Q3 Security Remediation Brief"  |  Due: "2026-10-24"  |  Confidence: 94%'
    cp2.font.bold = True
    cp2.font.size = Pt(12)
    cp2.font.color.rgb = TEXT_MAIN
    cp3 = cbtf.add_paragraph()
    cp3.text = 'Evidence Quote: "Audit findings require compliance brief by Oct 24 latest" (Chars 142-198) [PROVENANCE_CONFIRMED]'
    cp3.font.size = Pt(11)
    cp3.font.color.rgb = AMBER

    notes3 = s3.notes_slide.notes_text_frame
    notes3.text = (
        "ARCHITECTURE (30s):\n"
        "ActionPilot introduces a 5-step deterministic pipeline. We ingest audio via Web Speech API or documents, "
        "ground every single task to exact character offsets in the source text, feed it through our conflict matrix, "
        "enforce human signoff, and then dispatch to Google Calendar, Slack, or GitHub."
    )

    # ==================== SLIDE 4: 5 HACKATHON FEATURES ====================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "Feature Showcase", "5 Hackathon-Winning Core Capabilities",
               "Built, tested, and running live in production:")

    add_card(s4, Inches(0.8), Inches(2.5), Inches(3.6), Inches(2.1),
             "1. Live Voice Recording",
             ["Real-time meeting note dictation using Web Speech API with streaming waveform & instant extraction."],
             PRIMARY, "🎙️")
    add_card(s4, Inches(4.8), Inches(2.5), Inches(3.6), Inches(2.1),
             "2. 1-Click Demo Scenarios",
             ["3 pre-loaded real-world stress tests: Product Launch, SOC-2 Incident, and Executive Offsite."],
             PRIMARY, "⚡")
    add_card(s4, Inches(8.8), Inches(2.5), Inches(3.6), Inches(2.1),
             "3. Side-by-Side Highlighting",
             ["Document Inspector with glowing amber <mark> tags proving zero hallucination."],
             PRIMARY, "🔍")

    add_card(s4, Inches(0.8), Inches(4.8), Inches(5.6), Inches(2.0),
             "4. Interactive AI Conflict Resolver",
             ["Detects opposing deadlines side-by-side, provides context reasoning, and applies 1-click human resolution with audit trail."],
             AMBER, "⚔️")
    add_card(s4, Inches(6.8), Inches(4.8), Inches(5.6), Inches(2.0),
             "5. 1-Click Dispatch Hub + Dry-Run",
             ["Google Calendar URL prefill, Slack webhook dispatch, GitHub issue generator, plus a risk-free demo simulation toggle."],
             ACCENT, "🚀")

    notes4 = s4.notes_slide.notes_text_frame
    notes4.text = (
        "5 KILLER FEATURES (30s):\n"
        "For this hackathon, we built 5 production features:\n"
        "1. Live Voice Recording right in the browser.\n"
        "2. 1-Click Demo Scenarios to stress-test in 1 second.\n"
        "3. Side-by-Side radiant evidence highlighting.\n"
        "4. Interactive AI Conflict Resolver that catches conflicting deadlines.\n"
        "5. 1-Click Dispatch Hub with a dry-run simulation toggle."
    )

    # ==================== SLIDE 5: LIVE DEMO CHOREOGRAPHY ====================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "Demonstration", "Live Demo: 90-Second Winning Choreography",
               "How to present the live application to judges with zero friction:")

    add_card(s5, Inches(0.8), Inches(2.5), Inches(5.6), Inches(1.3),
             "Step 1: 1-Click Scenario Load",
             ["Select '⚡ Load Demo Scenario' -> 'Product Launch (Conflicting Deadlines)'. Parsed in 1s."],
             PRIMARY, "1️⃣")
    add_card(s5, Inches(0.8), Inches(4.0), Inches(5.6), Inches(1.3),
             "Step 2: Prove Zero Hallucination",
             ["Click 'View in Source Context'. Shows radiant amber highlight of the exact original sentence."],
             PRIMARY, "2️⃣")
    add_card(s5, Inches(0.8), Inches(5.5), Inches(5.6), Inches(1.3),
             "Step 3: Resolve Conflict & Dispatch",
             ["Click '⚡ Resolve with AI', pick Oct 20 freeze date, Approve, and dispatch to Google Calendar!"],
             ACCENT, "3️⃣")

    # Right side: Visual Representation
    conflict_demo = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.5), Inches(5.6), Inches(4.3))
    conflict_demo.fill.solid()
    conflict_demo.fill.fore_color.rgb = RGBColor(22, 28, 45)
    conflict_demo.line.color.rgb = AMBER
    conflict_demo.line.width = Pt(2)
    cd_tf = conflict_demo.text_frame
    cd_tf.word_wrap = True
    cdp0 = cd_tf.paragraphs[0]
    cdp0.text = "⚡ CONFLICT DETECTION CARD (LIVE UI)"
    cdp0.font.bold = True
    cdp0.font.size = Pt(14)
    cdp0.font.color.rgb = AMBER

    cdp1 = cd_tf.add_paragraph()
    cdp1.text = "Contradiction Detected:"
    cdp1.font.bold = True
    cdp1.font.size = Pt(13)
    cdp1.font.color.rgb = TEXT_MAIN
    cdp1.space_before = Pt(10)

    cdp2 = cd_tf.add_paragraph()
    cdp2.text = "• Priya (Product): 'Target launch on Nov 2'\n• DevOps Lead: 'Feature freeze required by Oct 20'"
    cdp2.font.size = Pt(12)
    cdp2.font.color.rgb = TEXT_MUTED
    cdp2.space_before = Pt(6)

    cdp3 = cd_tf.add_paragraph()
    cdp3.text = "AI Recommendation:"
    cdp3.font.bold = True
    cdp3.font.size = Pt(13)
    cdp3.font.color.rgb = PRIMARY
    cdp3.space_before = Pt(12)

    cdp4 = cd_tf.add_paragraph()
    cdp4.text = "'Adopt Oct 20 freeze to prevent unstable release. Target Nov 2 for public marketing rollout.'"
    cdp4.font.size = Pt(12)
    cdp4.font.color.rgb = RGBColor(147, 197, 253)
    cdp4.space_before = Pt(4)

    cdp5 = cd_tf.add_paragraph()
    cdp5.text = "Action: [Resolve with Oct 20] -> Logged to Audit Trail"
    cdp5.font.bold = True
    cdp5.font.size = Pt(12)
    cdp5.font.color.rgb = ACCENT
    cdp5.space_before = Pt(14)

    notes5 = s5.notes_slide.notes_text_frame
    notes5.text = (
        "LIVE DEMO CHOREOGRAPHY (90s):\n"
        "Let me show you this in action. We load our Product Launch sync. Notice how Priya asked for Nov 2, "
        "but DevOps mandated an Oct 20 freeze. ActionPilot immediately flagged this amber contradiction. "
        "With one click, the manager selects the safe freeze date, and dispatches the event straight to Google Calendar."
    )

    # ==================== SLIDE 6: DEEP TECH & SECURITY ====================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "Engineering Depth", "Production-Grade Stack & Enterprise Security",
               "Engineered with strict multi-tenant isolation and zero-downtime fail-safes:")

    add_card(s6, Inches(0.8), Inches(2.5), Inches(2.7), Inches(4.2),
             "Frontend",
             ["• React 18 + Vite\n• TypeScript\n• Web Speech API\n• Responsive Dark UI\n• Sub-second renders\n• Zero CORS proxy"],
             PRIMARY, "⚛️")
    add_card(s6, Inches(3.8), Inches(2.5), Inches(2.7), Inches(4.2),
             "Backend API",
             ["• FastAPI (Python 3.11)\n• High-throughput async ASGI routes\n• Auto-generated OpenAPI / Swagger docs\n• Pytest automated test suite"],
             PRIMARY, "🐍")
    add_card(s6, Inches(6.8), Inches(2.5), Inches(2.7), Inches(4.2),
             "Tenant Auth",
             ["• PBKDF2-SHA256 password hashing\n• Cryptographic JWT tokens\n• 6-digit Gmail OTP verification\n• Tenant-isolated SQLite schemas"],
             ACCENT, "🔐")
    add_card(s6, Inches(9.8), Inches(2.5), Inches(2.7), Inches(4.2),
             "Offline Fallback",
             ["• Dual-Engine Architecture\n• If external LLMs fail or go offline, local regex heuristics execute seamlessly with zero downtime"],
             AMBER, "🛡️")

    notes6 = s6.notes_slide.notes_text_frame
    notes6.text = (
        "DEEP TECH & SECURITY (30s):\n"
        "Under the hood, this isn't a simple wrapper. We have a FastAPI backend, tenant-isolated SQLite database, "
        "PBKDF2 password hashing, and Gmail OTP verification. Crucially, we built a dual-engine extractor — "
        "if external LLM APIs fail, our deterministic heuristic engine keeps running offline with zero downtime."
    )

    # ==================== SLIDE 7: ENTERPRISE ROI ====================
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_header(s7, "Business Impact", "Enterprise ROI & Market Opportunity",
               "Transforming operational overhead into deterministic velocity:")

    metrics = [
        ("92%", "Cycle Time Reduction", "From meeting conclusion to external dispatch across Calendar & Slack in under 60 seconds.", ACCENT),
        ("0%", "Hallucination Rate", "Strict character-level quote provenance prevents AI from inventing ungrounded actions.", PRIMARY),
        ("100%", "Audit Compliance", "Every resolution, edit, approval, and dispatch is cryptographically logged for SOC-2.", AMBER),
    ]

    for idx, (num, title, desc, col) in enumerate(metrics):
        x = Inches(0.8 + idx * 4.0)
        c = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.6), Inches(3.6), Inches(4.1))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = col
        c.line.width = Pt(1.5)
        ctf = c.text_frame
        ctf.word_wrap = True

        p0 = ctf.paragraphs[0]
        p0.text = num
        p0.font.bold = True
        p0.font.size = Pt(44)
        p0.font.color.rgb = col

        p1 = ctf.add_paragraph()
        p1.text = title
        p1.font.bold = True
        p1.font.size = Pt(18)
        p1.font.color.rgb = TEXT_MAIN
        p1.space_before = Pt(10)

        p2 = ctf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(12)

    notes7 = s7.notes_slide.notes_text_frame
    notes7.text = (
        "BUSINESS IMPACT (30s):\n"
        "The business impact is immediate: a 92% reduction in meeting-to-execution cycle time, "
        "zero hallucinations because of strict quote provenance, and 100% compliance readiness "
        "with cryptographic audit trails for every human approval."
    )

    # ==================== SLIDE 8: ROADMAP & LINKS ====================
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8)
    add_header(s8, "The Future & Q&A", "Roadmap & Live Verification Links",
               "Scaling ActionPilot into the enterprise operational nervous system:")

    add_card(s8, Inches(0.8), Inches(2.5), Inches(5.6), Inches(4.2),
             "Upcoming Product Milestones",
             ["• Q4 2026: Live Zoom, Google Meet & Microsoft Teams bot listener for continuous voice capture.",
              "• Q1 2027: Bidirectional 2-way sync with Jira, Linear, and Asana.",
              "• Q2 2027: Custom enterprise compliance policies & automated escalation matrix."],
             PRIMARY, "🚀")

    links_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.5), Inches(5.6), Inches(4.2))
    links_card.fill.solid()
    links_card.fill.fore_color.rgb = RGBColor(22, 28, 45)
    links_card.line.color.rgb = ACCENT
    links_card.line.width = Pt(1.5)
    ltf = links_card.text_frame
    ltf.word_wrap = True

    lp0 = ltf.paragraphs[0]
    lp0.text = "🔗 Live Deployment Links"
    lp0.font.bold = True
    lp0.font.size = Pt(18)
    lp0.font.color.rgb = ACCENT

    lp1 = ltf.add_paragraph()
    lp1.text = "• Backend API (Live on Render):\n  https://actionpilot-2oon.onrender.com/docs"
    lp1.font.size = Pt(12)
    lp1.font.color.rgb = TEXT_MAIN
    lp1.space_before = Pt(12)

    lp2 = ltf.add_paragraph()
    lp2.text = "• Open Source GitHub Repository:\n  https://github.com/satyamchaubey-gif/ActionPilot-"
    lp2.font.size = Pt(12)
    lp2.font.color.rgb = TEXT_MAIN
    lp2.space_before = Pt(12)

    lp3 = ltf.add_paragraph()
    lp3.text = "Thank you, Judges!\nWe are ready for your questions."
    lp3.font.bold = True
    lp3.font.size = Pt(16)
    lp3.font.color.rgb = AMBER
    lp3.space_before = Pt(20)

    notes8 = s8.notes_slide.notes_text_frame
    notes8.text = (
        "ROADMAP & CLOSING (15s):\n"
        "Our roadmap includes live Zoom bot listeners and bidirectional Jira sync. "
        "Our backend is already live on Render, and our repo is open source. "
        "Thank you, judges — we'd love to take your questions!"
    )

    prs.save(output_path)
    print(f"Presentation successfully created at {output_path}")

if __name__ == "__main__":
    create_deck()
